"""Table operation handlers — table add/remove."""

from __future__ import annotations

from openpyxl.worksheet.table import Table, TableStyleInfo

from fcp_core import OpResult, ParsedOp

from fcp_sheets.lib.table_styles import resolve_table_style
from fcp_sheets.server.resolvers import SheetsOpContext


def _table_add(op: ParsedOp, ctx: SheetsOpContext) -> OpResult:
    """table add NAME range:RANGE [style:STYLE] [banded-rows] [banded-cols] [first-col] [last-col]"""
    if len(op.positionals) < 2:
        return OpResult(
            success=False,
            message="Usage: table add NAME range:RANGE [style:STYLE] [banded-rows] [banded-cols]",
        )

    name = op.positionals[1]
    ws = ctx.active_sheet

    # Check for duplicate table name (ws.tables keys are display names)
    if name in ws.tables:
        return OpResult(success=False, message=f"Table '{name}' already exists")

    # Range is required
    range_str = op.params.get("range")
    if not range_str:
        return OpResult(success=False, message="Missing required param: range:RANGE")

    # Resolve style
    style_name = op.params.get("style", "TableStyleMedium9")
    try:
        style_name = resolve_table_style(style_name)
    except ValueError as e:
        return OpResult(success=False, message=str(e))

    # Boolean flags from positionals
    positional_set = set(p.lower() for p in op.positionals[2:])
    banded_rows = "banded-rows" in positional_set
    banded_cols = "banded-cols" in positional_set
    first_col = "first-col" in positional_set
    last_col = "last-col" in positional_set

    style_info = TableStyleInfo(
        name=style_name,
        showFirstColumn=first_col,
        showLastColumn=last_col,
        showRowStripes=banded_rows,
        showColumnStripes=banded_cols,
    )

    table = Table(displayName=name, ref=range_str, tableStyleInfo=style_info)
    ws.add_table(table)

    return OpResult(success=True, message=f"Table '{name}' added ({range_str})", prefix="+")


def _table_remove(op: ParsedOp, ctx: SheetsOpContext) -> OpResult:
    """table remove NAME"""
    if len(op.positionals) < 2:
        return OpResult(success=False, message="Usage: table remove NAME")

    name = op.positionals[1]
    ws = ctx.active_sheet

    # Find and remove table by display name (keys in ws.tables are display names)
    if name in ws.tables:
        del ws.tables[name]
        return OpResult(success=True, message=f"Table '{name}' removed", prefix="-")

    return OpResult(success=False, message=f"Table not found: {name!r}")


# Sub-command dispatch
_TABLE_SUBCMDS: dict[str, callable] = {
    "add": _table_add,
    "remove": _table_remove,
}


def op_table(op: ParsedOp, ctx: SheetsOpContext) -> OpResult:
    """Main table verb dispatcher."""
    if not op.positionals:
        return OpResult(success=False, message="Usage: table add|remove ...")

    subcmd = op.positionals[0].lower()
    handler = _TABLE_SUBCMDS.get(subcmd)
    if handler is None:
        available = ", ".join(sorted(_TABLE_SUBCMDS.keys()))
        return OpResult(success=False, message=f"Unknown table sub-command: {subcmd!r}. Available: {available}")

    return handler(op, ctx)


HANDLERS: dict[str, callable] = {
    "table": op_table,
}
